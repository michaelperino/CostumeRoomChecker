from bs4 import BeautifulSoup
from openpyxl import load_workbook
import os
# The lizard people won't tell you this, but you might also need html5lib and lxml installed. I'm unsure because
# Pandas made me do that when I was trying to do this earlier

# My pandas solution failed, but ChatGPT wrote the following:
# Updated to read from local HTML files instead of making web requests

DRESS_URL = "https://oldschool.runescape.wiki/w/Mahogany_fancy_dress_box"
ARMOUR_URL = "https://oldschool.runescape.wiki/w/Mahogany_armour_case"
CAPE_URL = "https://oldschool.runescape.wiki/w/Magical_cape_rack"
WARDROBE_URL = "https://oldschool.runescape.wiki/w/Marble_magic_wardrobe"
TOY_URL = "https://oldschool.runescape.wiki/w/Mahogany_toy_box"
TREASURE_URL = "https://oldschool.runescape.wiki/w/Mahogany_treasure_chest"

# Map URLs to local HTML files
html_files = {
    DRESS_URL: "html_files/Mahogany fancy dress box - OSRS Wiki.htm",
    ARMOUR_URL: "html_files/Mahogany armour case - OSRS Wiki.htm",
    CAPE_URL: "html_files/Magical cape rack - OSRS Wiki.htm",
    WARDROBE_URL: "html_files/Marble magic wardrobe - OSRS Wiki.htm",
    TOY_URL: "html_files/Mahogany toy box - OSRS Wiki.htm",
    TREASURE_URL: "html_files/Mahogany treasure chest - OSRS Wiki.htm",
}


# Derive human-readable page titles from the local HTML filenames
def _page_title_from_path(path: str) -> str:
    name = os.path.basename(path)
    # Strip the suffix " - OSRS Wiki.htm"
    if name.endswith(" - OSRS Wiki.htm"):
        name = name[: -len(" - OSRS Wiki.htm")]
    return name


html_page_titles = {
    url: _page_title_from_path(path) for url, path in html_files.items()
}


# Consistent normalization for matching names across sources
def sanitize_name(s: str) -> str:
    if not isinstance(s, str):
        s = str(s)
    return (
        s.strip()
        .upper()
        .replace(" ", "_")
        .replace("%27", "")
        .replace("'", "")
        .replace("(", "")
        .replace(")", "")
    )


urls = [DRESS_URL, ARMOUR_URL, CAPE_URL, WARDROBE_URL, TOY_URL, TREASURE_URL]
page_items = {}

with open("OUT_WIKI_LINKS.txt", "w") as f:
    for url in urls:
        f.write("\n\n")
        f.write(str(url))
        f.write("\n")
        # Read from local HTML file instead of making a web request
        with open(html_files[url], "r", encoding="utf-8") as html_file:
            html = html_file.read()
        soup = BeautifulSoup(html, "html.parser")

        # Find all tables in the page
        tables = soup.find_all("table")

        # Filter out tables with the class "navbox"
        filtered_tables = [
            table
            for table in tables
            if "navbox" not in table.get("class", [])
            and "navbox-subgroup" not in table.get("class", [])
        ]

        # Check if there are any filtered tables
        if not filtered_tables:
            raise ValueError("No tables found that do not have the class 'navbox'.")

        filtered_tables_2 = [filtered_tables[-1]]
        if url == TREASURE_URL:
            filtered_tables_2 = filtered_tables[-6:]

        # Collect items for this page
        page_key = sanitize_name(html_page_titles[url])
        items_set = page_items.setdefault(page_key, set())

        for target_table in filtered_tables_2:
            # Extract hyperlinks from the second column of each row in the table
            rows = target_table.find_all("tr")
            for row in rows:
                columns = row.find_all("td")
                if len(columns) >= 2:
                    second_column = columns[1]
                    if url == TOY_URL:
                        second_column = columns[0]
                    for a in second_column.find_all("a"):
                        href = a.get("href", "")
                        upper_href = href.upper()
                        if upper_href.startswith("/W/"):
                            page = href[3:]
                        else:
                            idx = upper_href.find("/W/")
                            page = href[idx + 3 :] if idx != -1 else href
                        cleaned = sanitize_name(page)
                        if cleaned:
                            items_set.add(cleaned)
                            f.write(cleaned + "\n")

# Read back in the file, and if a line contains /W/, only print after the /W/
with open("OUT_WIKI_LINKS.txt", "r") as f:
    lines = f.readlines()

BLACKLIST = ["TE_NOTE-TROUVER-1"]
with open("OUT_WIKI_LINKS_CLEANED.txt", "w") as f:
    for line in lines:
        if line.strip() in BLACKLIST:
            print("BLACKLIST1", line.strip())
            continue
        if "/W/" in line:
            cleaned_line = line.split("/W/")[-1]
            if cleaned_line.strip() not in BLACKLIST:
                f.write(cleaned_line)
            else:
                print("BLACKLIST2", cleaned_line.strip())
        else:
            f.write(line)

# ===== Excel comparison and missing lists generation =====
# Build Excel-derived sets for column A per sheet, then diff against wiki page items


def load_excel_column_a_sets(xlsm_path: str):
    wb = load_workbook(xlsm_path, read_only=True, data_only=True)
    sheet_to_items = {}
    for wsname in wb.sheetnames:
        ws = wb[wsname]
        excel_items = set()
        # Read only column A via iter_rows (works in read_only mode)
        for row in ws.iter_rows(min_row=1, max_col=1, values_only=True):
            val = row[0]
            if val is None:
                continue
            cleaned = sanitize_name(val)
            if cleaned:
                excel_items.add(cleaned)
        sheet_to_items[sanitize_name(wsname)] = {
            "raw_name": wsname,
            "items": excel_items,
        }
    return sheet_to_items


def write_missing_lists(
    sheet_sets: dict, wiki_sets: dict, out_dir: str = "MISSING_LISTS"
):
    os.makedirs(out_dir, exist_ok=True)
    summary_lines = []

    for sheet_key, payload in sheet_sets.items():
        sheet_raw = payload["raw_name"]
        excel_items = payload["items"]

        wiki_items = wiki_sets.get(sheet_key, set())
        missing = sorted(wiki_items - excel_items)

        out_path = os.path.join(out_dir, f"MISSING_{sanitize_name(sheet_raw)}.txt")
        with open(out_path, "w", encoding="utf-8") as f:
            for item in missing:
                f.write(item + "\n")

        intersect_count = len(wiki_items & excel_items)
        summary_lines.append(
            f"{sheet_raw}: wiki={len(wiki_items)} excel={len(excel_items)} missing={len(missing)} intersect={intersect_count}"
        )

    with open(os.path.join(out_dir, "MISSING_SUMMARY.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(summary_lines) + "\n")


# Load Excel and produce missing lists
excel_sets = load_excel_column_a_sets("WikiParser.xlsm")
write_missing_lists(excel_sets, page_items)

# Debug: write out raw sets for inspection
_debug_dir = os.path.join("MISSING_LISTS", "DEBUG")
os.makedirs(_debug_dir, exist_ok=True)
for _sheet_key, _payload in excel_sets.items():
    _wiki_items = page_items.get(_sheet_key, set())
    _excel_items = _payload["items"]
    if _wiki_items:
        _safe = sanitize_name(_payload["raw_name"])
        with open(
            os.path.join(_debug_dir, f"EXCEL_{_safe}.txt"), "w", encoding="utf-8"
        ) as f:
            for item in sorted(_excel_items):
                f.write(item + "\n")
        with open(
            os.path.join(_debug_dir, f"WIKI_{_safe}.txt"), "w", encoding="utf-8"
        ) as f:
            for item in sorted(_wiki_items):
                f.write(item + "\n")
